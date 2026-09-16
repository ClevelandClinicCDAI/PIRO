import asyncio
import csv
import io
import json
from types import SimpleNamespace

from openpyxl import load_workbook

from apis.version1 import route_extraction
from core.constants import Constants


class _FakeQueueItem:
    def __init__(self, case_id: int, case_number: str):
        self.ExtractionQueueId = case_id
        self.ExtractionSessionId = 7
        self.CaseId = case_id
        self.Status = "pending"
        self.ErrorMessage = None
        self.AttemptCount = 0
        self.CreateDate = None
        self.Case = SimpleNamespace(CaseNumber=case_number)

    @property
    def CaseNumber(self):
        return self.Case.CaseNumber


class _FakeResult:
    def __init__(
        self,
        case_id: int,
        case_number: str,
        field_name: str,
        extracted_value: str,
    ):
        self.ExtractionResultId = case_id
        self.ExtractionRunId = 9
        self.ExtractionSessionId = 7
        self.CaseId = case_id
        self.Case = SimpleNamespace(CaseNumber=case_number)
        self.FieldName = field_name
        self.ExtractedValue = extracted_value
        self.ReviewedValue = None
        self.Confidence = 0.9
        self.ProvenanceText = None
        self.SourceCommentId = None
        self.IsReviewed = False
        self.IsIncorrect = False
        self.ReviewedBy = None
        self.ReviewedDate = None

    @property
    def CaseNumber(self):
        return self.Case.CaseNumber


async def _read_stream(response) -> bytes:
    return b"".join([chunk async for chunk in response.body_iterator])


def test_add_to_queue_masks_case_numbers_for_demoadmin(monkeypatch):
    queue_items = [_FakeQueueItem(11, "REAL-123")]

    monkeypatch.setattr(
        route_extraction,
        "_require_session_ownership",
        lambda *args, **kwargs: SimpleNamespace(),
    )
    monkeypatch.setattr(
        route_extraction, "add_cases_to_queue", lambda **kwargs: None
    )
    monkeypatch.setattr(
        route_extraction, "get_queue", lambda session_id, db: queue_items
    )

    result = asyncio.run(
        route_extraction.add_to_queue(
            route_extraction.ExtractionQueueAdd(session_id=7, case_ids=[11]),
            current_user="demo@example.com",
            current_user_id=1,
            current_role=Constants.RoleDemoAdmin,
            db=SimpleNamespace(),
        )
    )

    assert result[0].CaseNumber == "-"
    assert queue_items[0].Case.CaseNumber == "REAL-123"


def test_add_saved_search_to_queue_masks_case_numbers_for_demoadmin(
    monkeypatch,
):
    import core.search_util as search_util
    import solr.repository.piro as solr_repo

    queue_items = [_FakeQueueItem(12, "REAL-456")]
    saved_search = SimpleNamespace(
        SearchQuery="https://example.test?searchFilter=[]",
        AdvancedQuery=None,
        MRN=None,
    )

    monkeypatch.setattr(
        route_extraction,
        "_require_session_ownership",
        lambda *args, **kwargs: SimpleNamespace(),
    )
    monkeypatch.setattr(
        route_extraction, "get_search", lambda *args: saved_search
    )
    monkeypatch.setattr(
        route_extraction, "add_cases_to_queue", lambda **kwargs: None
    )
    monkeypatch.setattr(
        route_extraction, "get_queue", lambda session_id, db: queue_items
    )
    monkeypatch.setattr(search_util, "filter_str_object", lambda raw: {})
    monkeypatch.setattr(
        solr_repo,
        "search_Q",
        lambda **kwargs: {"caseIds": [12], "total": 1},
    )

    result = asyncio.run(
        route_extraction.add_saved_search_to_queue(
            route_extraction.ExtractionQueueFromSearch(
                session_id=7,
                search_id=99,
            ),
            current_user="demo@example.com",
            current_user_id=1,
            current_role=Constants.RoleDemoAdmin,
            db=SimpleNamespace(),
            solr=SimpleNamespace(),
        )
    )

    assert result[0].CaseNumber == "-"
    assert queue_items[0].Case.CaseNumber == "REAL-456"


def test_export_results_masks_demoadmin_case_numbers_without_merging_rows(
    monkeypatch,
):
    session = SimpleNamespace(
        SchemaJson='{"field_a": {"type": "string"}}',
        Name="Demo Session",
    )
    results = [
        _FakeResult(11, "REAL-123", "field_a", '"value a"'),
        _FakeResult(12, "REAL-456", "field_a", '"value b"'),
    ]

    monkeypatch.setattr(
        route_extraction,
        "_require_session_read_access",
        lambda *args, **kwargs: session,
    )
    monkeypatch.setattr(
        route_extraction,
        "get_results_for_session",
        lambda session_id, db: results,
    )

    json_response = asyncio.run(
        route_extraction.export_results(
            7,
            format="json",
            current_user_id=1,
            current_role=Constants.RoleDemoAdmin,
            db=SimpleNamespace(),
        )
    )
    json_rows = json.loads(json_response.body)
    assert len(json_rows) == 2
    assert [row["case_number"] for row in json_rows] == ["-", "-"]
    assert [row["field_a"] for row in json_rows] == ["value a", "value b"]

    csv_response = asyncio.run(
        route_extraction.export_results(
            7,
            format="csv",
            current_user_id=1,
            current_role=Constants.RoleDemoAdmin,
            db=SimpleNamespace(),
        )
    )
    csv_rows = list(
        csv.DictReader(
            io.StringIO(
                asyncio.run(_read_stream(csv_response)).decode("utf-8")
            )
        )
    )
    assert len(csv_rows) == 2
    assert [row["case_number"] for row in csv_rows] == ["-", "-"]

    excel_response = asyncio.run(
        route_extraction.export_results(
            7,
            format="excel",
            current_user_id=1,
            current_role=Constants.RoleDemoAdmin,
            db=SimpleNamespace(),
        )
    )
    workbook = load_workbook(
        io.BytesIO(asyncio.run(_read_stream(excel_response)))
    )
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet.cell(row=2, column=1).value == "-"
    assert worksheet.cell(row=3, column=1).value == "-"
    assert worksheet.cell(row=2, column=2).value == "value a"
    assert worksheet.cell(row=3, column=2).value == "value b"
